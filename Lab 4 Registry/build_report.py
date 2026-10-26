# -*- coding: utf-8 -*-
"""Собрать otchet4_Registry.docx: код с полужирными именами API, как в Лаб4_СП_Отчет.docx."""
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, Twips
from lxml import etree

ROOT = Path(__file__).resolve().parent
MAIN = ROOT / "main.c"
OUT = ROOT / "otchet4_Registry.docx"

# Имена функций и вспомогательных процедур — полужирным, по образцу отчёта
BOLD_RE = re.compile(
    r"\b(?:"
    r"Reg[A-Za-z]+|"
    r"CreateEventW|WaitForSingleObject|CloseHandle|"
    r"GetEnvironmentVariableW|GetFileAttributesExW|GetLastError|"
    r"GetSecurityDescriptorDacl|GetAclInformation|GetAce|"
    r"malloc|free|fflush|system|_setmode|_fileno|"
    r"wcslen|swprintf|wprintf|"
    r"PrintResult|PrintMegabytes|PrintRegistrySize|PrintMask|PrintAccess"
    r")\b"
)


def set_font(run, name, size, bold=False):
    run.bold = bold
    run.font.name = name
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = etree.SubElement(rpr, qn("w:rFonts"))
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)


def add_code_line(doc, line):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    text = line if line else " "
    pos = 0
    for match in BOLD_RE.finditer(text):
        if match.start() > pos:
            run = p.add_run(text[pos : match.start()])
            set_font(run, "Consolas", 8.5, False)
        run = p.add_run(match.group())
        set_font(run, "Consolas", 8.5, True)
        pos = match.end()
    if pos < len(text):
        run = p.add_run(text[pos:])
        set_font(run, "Consolas", 8.5, False)


def main():
    lines = MAIN.read_text(encoding="utf-8").splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Twips(11906)
    sec.page_height = Twips(16838)
    for margin in ("top_margin", "right_margin", "bottom_margin", "left_margin"):
        setattr(sec, margin, Twips(1440))
    sec.header_distance = Twips(720)
    sec.footer_distance = Twips(720)

    body = doc.element.body
    for child in list(body):
        if child.tag == qn("w:p"):
            body.remove(child)

    p = doc.add_paragraph()
    run = p.add_run("ФИО группа")
    set_font(run, "Times New Roman", 13, False)
    add_code_line(doc, "")

    for line in lines:
        add_code_line(doc, line)

    doc.save(OUT)
    print("saved", OUT, "lines", len(lines))


if __name__ == "__main__":
    main()
